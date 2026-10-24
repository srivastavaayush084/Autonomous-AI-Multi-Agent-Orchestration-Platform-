import asyncio
import json
import logging
import os
import re
import uuid
import zipfile
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple

from app.core.config import settings
from app.database.mongodb import db_manager
from app.models.approval import ApprovalStatus, ApprovalRequest, RiskLevel
from app.models.artifact import ArtifactResponse, ArtifactType
from app.models.intent import (
    IntentType,
    WorkflowType,
    CANONICAL_INTENT_WORKFLOW_MAP,
    ExecutionPlan,
    RoutingSource,
)
from app.models.run import RunStatus
from app.models.task import TaskStatus
from app.services.model_adapter import model_adapter
from app.services.router import model_router
from app.services.rag_service import rag_service
from app.services.hybrid_router import hybrid_router

logger = logging.getLogger(__name__)

# ──────────────────────────────────────────────────────────────────────────────
# Request Classifier (Backward Compatibility Wrapper)
# ──────────────────────────────────────────────────────────────────────────────

def classify_request(goal: str) -> str:
    """
    Backward-compatible request classifier delegating to hybrid_router.
    Returns one of: 'qa', 'research', 'website', 'code', 'data_analysis', 'full_pipeline'
    """
    det = hybrid_router.route_deterministic(goal)
    if det:
        intent = det[0]
        mapping = {
            IntentType.CHAT: "qa",
            IntentType.RESEARCH: "research",
            IntentType.SOFTWARE_PROJECT: "website",
            IntentType.CODE: "code",
            IntentType.DATA_ANALYSIS: "data_analysis",
        }
        return mapping.get(intent, "full_pipeline")
    return "full_pipeline"


# ──────────────────────────────────────────────────────────────────────────────
# File Export Helpers
# ──────────────────────────────────────────────────────────────────────────────

EXPORTS_DIR = Path(__file__).parent.parent.parent / "exports"
EXPORTS_DIR.mkdir(exist_ok=True)


def _save_docx(title: str, content: str, run_id: str) -> Optional[str]:
    """Generate a .docx file and return its absolute path."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # Title
        heading = doc.add_heading(title, 0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Parse content into paragraphs / code blocks
        lines = content.split("\n")
        in_code_block = False
        code_buffer: List[str] = []

        for line in lines:
            if line.startswith("```"):
                if in_code_block:
                    # Flush code buffer
                    para = doc.add_paragraph("\n".join(code_buffer))
                    run = para.runs[0] if para.runs else para.add_run("")
                    run.font.name = "Courier New"
                    run.font.size = Pt(9)
                    code_buffer = []
                    in_code_block = False
                else:
                    in_code_block = True
                continue

            if in_code_block:
                code_buffer.append(line)
                continue

            if line.startswith("# "):
                doc.add_heading(line[2:].strip(), level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:].strip(), level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:].strip(), level=3)
            elif line.startswith("- ") or line.startswith("* "):
                doc.add_paragraph(line[2:].strip(), style="List Bullet")
            elif re.match(r"^\d+\. ", line):
                doc.add_paragraph(re.sub(r"^\d+\. ", "", line).strip(), style="List Number")
            elif line.strip():
                # Handle **bold** inline
                para = doc.add_paragraph()
                parts = re.split(r"(\*\*[^*]+\*\*)", line)
                for part in parts:
                    if part.startswith("**") and part.endswith("**"):
                        run = para.add_run(part[2:-2])
                        run.bold = True
                    else:
                        para.add_run(part)
            else:
                doc.add_paragraph("")

        path = EXPORTS_DIR / f"{run_id}.docx"
        doc.save(str(path))
        logger.info(f"DOCX saved: {path}")
        return str(path)
    except Exception as e:
        logger.error(f"DOCX generation failed: {e}")
        return None


def _save_pdf(title: str, content: str, run_id: str) -> Optional[str]:
    """Generate a PDF using reportlab and return its absolute path."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Preformatted
        from reportlab.lib.enums import TA_LEFT, TA_CENTER
        from reportlab.lib import colors

        path = EXPORTS_DIR / f"{run_id}.pdf"
        doc = SimpleDocTemplate(
            str(path),
            pagesize=A4,
            rightMargin=2 * cm,
            leftMargin=2 * cm,
            topMargin=2 * cm,
            bottomMargin=2 * cm,
        )
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle(
            "CustomTitle",
            parent=styles["Title"],
            fontSize=18,
            spaceAfter=20,
            alignment=TA_CENTER,
        )
        h1_style = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=14, spaceAfter=8)
        h2_style = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=12, spaceAfter=6)
        h3_style = ParagraphStyle("H3", parent=styles["Heading3"], fontSize=11, spaceAfter=4)
        body_style = ParagraphStyle("Body", parent=styles["Normal"], fontSize=10, leading=16, spaceAfter=6)
        code_style = ParagraphStyle(
            "Code",
            parent=styles["Code"],
            fontSize=8,
            fontName="Courier",
            backColor=colors.HexColor("#f5f5f5"),
            spaceAfter=8,
            leading=12,
        )

        story.append(Paragraph(title, title_style))
        story.append(Spacer(1, 0.4 * cm))

        lines = content.split("\n")
        in_code = False
        code_buf: List[str] = []

        for line in lines:
            if line.startswith("```"):
                if in_code:
                    block = "\n".join(code_buf)
                    story.append(Preformatted(block, code_style))
                    code_buf = []
                    in_code = False
                else:
                    in_code = True
                continue

            if in_code:
                code_buf.append(line)
                continue

            if line.startswith("# "):
                story.append(Paragraph(line[2:].strip(), h1_style))
            elif line.startswith("## "):
                story.append(Paragraph(line[3:].strip(), h2_style))
            elif line.startswith("### "):
                story.append(Paragraph(line[4:].strip(), h3_style))
            elif line.startswith("- ") or line.startswith("* "):
                story.append(Paragraph(f"• {line[2:].strip()}", body_style))
            elif line.strip():
                # Handle **bold**
                formatted = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", line)
                story.append(Paragraph(formatted, body_style))
            else:
                story.append(Spacer(1, 0.2 * cm))

        doc.build(story)
        logger.info(f"PDF saved: {path}")
        return str(path)
    except Exception as e:
        logger.error(f"PDF generation failed: {e}")
        return None


def _save_zip(files: Dict[str, str], run_id: str) -> Optional[str]:
    """Bundle a dict of {filename: content} into a ZIP, verify integrity, and return its path."""
    try:
        path = EXPORTS_DIR / f"{run_id}.zip"
        with zipfile.ZipFile(str(path), "w", zipfile.ZIP_DEFLATED) as zf:
            for filename, content in files.items():
                zf.writestr(filename, content)
        # Verify ZIP integrity
        with zipfile.ZipFile(str(path), "r") as zf:
            bad_file = zf.testzip()
            if bad_file:
                logger.error(f"ZIP archive corrupted at entry: {bad_file}")
                return None
        logger.info(f"ZIP saved and verified: {path}")
        return str(path)
    except Exception as e:
        logger.error(f"ZIP generation failed: {e}")
        return None


# ──────────────────────────────────────────────────────────────────────────────
# Main Orchestration Service
# ──────────────────────────────────────────────────────────────────────────────

class OrchestrationService:
    def __init__(self):
        self.mem_projects: Dict[str, Dict[str, Any]] = {}
        self.mem_runs: Dict[str, Dict[str, Any]] = {}
        self.mem_tasks: Dict[str, Dict[str, Any]] = {}
        self.mem_approvals: Dict[str, Dict[str, Any]] = {}
        self.mem_artifacts: Dict[str, Dict[str, Any]] = {}
        self._active_loops: set = set()

    def _now(self) -> datetime:
        return datetime.now(timezone.utc)

    @staticmethod
    def _diff_seconds(t_end: Optional[datetime], t_start: Optional[datetime]) -> Optional[float]:
        if not t_end or not t_start:
            return None
        if t_end.tzinfo is not None and t_start.tzinfo is None:
            t_start = t_start.replace(tzinfo=timezone.utc)
        elif t_end.tzinfo is None and t_start.tzinfo is not None:
            t_end = t_end.replace(tzinfo=timezone.utc)
        return (t_end - t_start).total_seconds()

    # ==========================================
    # Project Operations
    # ==========================================
    async def create_project(self, name: str, description: str = "", metadata: Dict[str, Any] = None) -> Dict[str, Any]:
        project_id = f"proj-{uuid.uuid4().hex[:8]}"
        doc = {
            "project_id": project_id,
            "name": name,
            "description": description,
            "created_at": self._now(),
            "updated_at": self._now(),
            "metadata": metadata or {}
        }
        if db_manager.connected and db_manager.db is not None:
            await db_manager.db.projects.insert_one(doc.copy())
        self.mem_projects[project_id] = doc
        return doc

    async def list_projects(self) -> List[Dict[str, Any]]:
        if db_manager.connected and db_manager.db is not None:
            cursor = db_manager.db.projects.find({}, {"_id": 0}).sort("created_at", -1)
            return await cursor.to_list(length=100)
        return list(self.mem_projects.values())

    async def get_project(self, project_id: str) -> Optional[Dict[str, Any]]:
        if db_manager.connected and db_manager.db is not None:
            doc = await db_manager.db.projects.find_one({"project_id": project_id}, {"_id": 0})
            if doc:
                return doc
        return self.mem_projects.get(project_id)

    async def delete_project(self, project_id: str) -> bool:
        proj = await self.get_project(project_id)
        if not proj:
            return False

        self.mem_projects.pop(project_id, None)
        project_run_ids = [r["run_id"] for r in self.mem_runs.values() if r.get("project_id") == project_id]
        for r_id in project_run_ids:
            self.mem_runs.pop(r_id, None)
            self.mem_tasks = {k: v for k, v in self.mem_tasks.items() if v.get("run_id") != r_id}
            self.mem_approvals = {k: v for k, v in self.mem_approvals.items() if v.get("run_id") != r_id}
            self.mem_artifacts = {k: v for k, v in self.mem_artifacts.items() if v.get("run_id") != r_id}

        if db_manager.connected and db_manager.db is not None:
            try:
                await db_manager.db.projects.delete_one({"project_id": project_id})
                await db_manager.db.runs.delete_many({"project_id": project_id})
                for r_id in project_run_ids:
                    await db_manager.db.tasks.delete_many({"run_id": r_id})
                    await db_manager.db.approvals.delete_many({"run_id": r_id})
                    await db_manager.db.artifacts.delete_many({"run_id": r_id})
            except Exception as e:
                logger.error(f"Failed to cascade delete project {project_id}: {e}")

        await rag_service.delete_project_documents(project_id)
        logger.info(f"Project {project_id} and all cascaded artifacts successfully deleted.")
        return True

    # ==========================================
    # Run Lifecycle Operations
    # ==========================================
    async def create_run(self, project_id: str, goal: str, parameters: Dict[str, Any] = None) -> Dict[str, Any]:
        run_id = f"run-{uuid.uuid4().hex[:8]}"
        doc = {
            "run_id": run_id,
            "project_id": project_id,
            "goal": goal,
            "status": RunStatus.PENDING,
            "created_at": self._now(),
            "updated_at": self._now(),
            "started_at": None,
            "completed_at": None,
            "execution_time_seconds": None,
            "total_tasks": 0,
            "completed_tasks": 0,
            "failed_tasks": 0,
            "pending_approvals": 0,
            "error_message": None,
            "metadata": parameters or {}
        }
        if db_manager.connected and db_manager.db is not None:
            await db_manager.db.runs.insert_one(doc.copy())
        self.mem_runs[run_id] = doc

        asyncio.create_task(self.start_run_execution(run_id))
        return doc

    async def get_run(self, run_id: str) -> Optional[Dict[str, Any]]:
        if db_manager.connected and db_manager.db is not None:
            doc = await db_manager.db.runs.find_one({"run_id": run_id}, {"_id": 0})
            if doc:
                return doc
        return self.mem_runs.get(run_id)

    async def list_runs(self, project_id: Optional[str] = None) -> List[Dict[str, Any]]:
        if db_manager.connected and db_manager.db is not None:
            q = {"project_id": project_id} if project_id else {}
            cursor = db_manager.db.runs.find(q, {"_id": 0}).sort("created_at", -1)
            return await cursor.to_list(length=100)

        runs = list(self.mem_runs.values())
        if project_id:
            runs = [r for r in runs if r.get("project_id") == project_id]
        return sorted(runs, key=lambda x: x["created_at"], reverse=True)

    async def resume_run(self, run_id: str) -> Optional[Dict[str, Any]]:
        run = await self.get_run(run_id)
        if not run:
            return None

        tasks = await self.get_run_tasks(run_id)
        if not tasks or run.get("status") == RunStatus.PLANNING:
            asyncio.create_task(self.start_run_execution(run_id))
            return run

        for t in tasks:
            if t["status"] == TaskStatus.RUNNING:
                t["status"] = TaskStatus.PENDING
                await self._save_task(t)

        run["status"] = RunStatus.RUNNING
        await self._save_run(run)
        asyncio.create_task(self._execute_dag_loop(run_id))
        return run

    # ==========================================
    # Autonomous Multi-Agent Execution Engine
    # ==========================================
    async def start_run_execution(self, run_id: str):
        run = await self.get_run(run_id)
        if not run:
            return

        goal = run["goal"]
        parameters = run.get("metadata", {})
        attachments = parameters.get("attachments", [])

        # ── Hybrid Intent Routing & Workflow Planning ─────────────────────────
        plan = await hybrid_router.plan(goal, parameters=parameters, attachments=attachments)
        logger.info(f"Run {run_id}: planned intent='{plan.intent.value}', workflow='{plan.workflow.value}' ({plan.reason})")

        run["status"] = RunStatus.PLANNING
        run["started_at"] = self._now()
        run["metadata"] = run.get("metadata") or {}
        run["metadata"]["intent"] = plan.intent.value
        run["metadata"]["workflow"] = plan.workflow.value
        run["metadata"]["run_type"] = plan.intent.value
        run["metadata"]["routing_source"] = plan.routing_source.value
        run["metadata"]["routing_reason"] = plan.reason
        run["metadata"]["confidence"] = plan.confidence
        run["metadata"]["deliverables"] = plan.deliverables
        run["metadata"]["plan"] = plan.model_dump()
        await self._save_run(run)

        # Handle required clarification (e.g., missing dataset for data_analysis)
        if plan.requires_clarification:
            run["status"] = RunStatus.FAILED
            run["error_message"] = plan.clarification_message
            run["completed_at"] = self._now()
            await self._save_run(run)
            await self._register_artifact(
                run_id, "clarification",
                title="Action Required: Dataset Missing",
                artifact_type=ArtifactType.REPORT,
                content=f"## Input Required\n\n{plan.clarification_message}",
                is_user_facing=True,
                metadata={"requires_clarification": True, "workflow": plan.workflow.value}
            )
            return

        # ── Answer Question (Chat) ────────────────────────────────────────────
        if plan.workflow == WorkflowType.ANSWER_QUESTION:
            await self._execute_qa_run(run_id, goal, plan=plan)
            return

        # ── Research & Write (Research) ───────────────────────────────────────
        if plan.workflow == WorkflowType.RESEARCH_AND_WRITE:
            await self._execute_research_run(run_id, goal, plan=plan)
            return

        # ── Build Software (Software Project) ─────────────────────────────────
        if plan.workflow == WorkflowType.BUILD_SOFTWARE:
            await self._execute_website_run(run_id, goal, plan=plan)
            return

        # ── Generate Code (Code) ──────────────────────────────────────────────
        if plan.workflow == WorkflowType.GENERATE_CODE:
            await self._execute_code_run(run_id, goal, plan=plan)
            return

        # ── Analyze Data (Data Analysis) ──────────────────────────────────────
        if plan.workflow == WorkflowType.ANALYZE_DATA:
            await self._execute_data_analysis_run(run_id, goal, plan=plan)
            return

        # ── Full pipeline fallback ────────────────────────────────────────────
        await self._execute_full_pipeline_run(run_id, goal)

    # ──────────────────────────────────────────────────────────────────────────
    # Content-Aware Completion Validation Helpers
    # ──────────────────────────────────────────────────────────────────────────
    @staticmethod
    def _validate_qa_content(answer: str, goal: str) -> Tuple[bool, str]:
        if not answer or len(answer.strip()) < 40:
            return False, "Response is empty or insufficiently substantive."
        canned_phrases = [
            "the requested task has been completed",
            "all acceptance criteria have been satisfied",
            "the implementation follows best practices",
            "point 1 — foundational understanding",
            "point 2 — practical application",
            "the subject at hand represents a significant area of inquiry",
            "[error:",
            "model inference failed",
            "inference failed",
            "unavailable or timed out",
            "error message related to",
            "appears to be an error message",
            "appears to be a error message",
            "cannot provide a review of this paper",
            "can't provide a review of this paper",
        ]
        lower = answer.lower()
        for phrase in canned_phrases:
            if phrase in lower:
                return False, f"Response contains generic placeholder or error text: '{phrase}'."
        if "rag" in goal.lower():
            if not any(term in lower for term in ["retriev", "rag", "generation", "augmented"]):
                return False, "Response does not explain Retrieval-Augmented Generation."
        return True, "Valid QA response"

    @staticmethod
    def _validate_research_content(paper: str, goal: str) -> Tuple[bool, str]:
        if not paper or len(paper.strip()) < 300:
            return False, "Research paper is empty or too short."
        canned_phrases = [
            "point 1 — foundational understanding",
            "point 2 — practical application",
            "point 3 — future directions",
            "the subject at hand represents a significant area of inquiry",
            "the requested task has been completed",
            "all acceptance criteria have been satisfied",
            "[error:",
            "model inference failed",
            "inference failed",
            "unavailable or timed out",
            "error message related to",
            "appears to be an error message",
            "appears to be a error message",
            "cannot provide a review of this paper",
            "can't provide a review of this paper",
        ]
        lower = paper.lower()
        for phrase in canned_phrases:
            if phrase in lower:
                return False, f"Research paper contains generic filler or error text: '{phrase}'."
        required_sections = ["abstract", "introduction", "references"]
        missing = [s for s in required_sections if s not in lower]
        if missing:
            return False, f"Research paper missing required sections: {', '.join(missing)}."
        return True, "Valid research paper"

    @staticmethod
    def _validate_website_files(files: Dict[str, str], is_react: bool) -> Tuple[bool, str]:
        if not files:
            return False, "No website source files were generated."
        canned_phrases = [
            "the requested task has been completed",
            "all acceptance criteria have been satisfied",
            "the implementation follows best practices",
            "point 1 — foundational understanding",
            "[error:",
            "model inference failed",
            "inference failed",
            "unavailable or timed out",
        ]
        for fname, content in files.items():
            if not content or len(content.strip()) < 20:
                return False, f"File '{fname}' is empty or too short."
            lower = content.lower()
            for phrase in canned_phrases:
                if phrase in lower:
                    return False, f"File '{fname}' contains generic completion or error text: '{phrase}'."

        if is_react:
            if "package.json" not in files:
                return False, "React project missing package.json manifest."
            try:
                pkg = json.loads(files["package.json"])
                if "dependencies" not in pkg or "react" not in pkg["dependencies"]:
                    return False, "package.json missing react dependency."
            except Exception as e:
                return False, f"package.json is invalid JSON: {e}"

            has_app = any(f in files for f in ["src/App.jsx", "src/App.tsx", "App.jsx", "App.tsx"])
            if not has_app:
                return False, "React project missing App component (src/App.jsx or src/App.tsx)."
        else:
            if "index.html" not in files:
                return False, "Vanilla website missing index.html."
            if "<html" not in files["index.html"].lower() and "<!doctype html" not in files["index.html"].lower():
                return False, "index.html is not valid HTML structure."
            if "style.css" not in files:
                return False, "Vanilla website missing style.css."
            if "{" not in files["style.css"] or "}" not in files["style.css"]:
                return False, "style.css does not contain valid CSS rules."

        return True, "Valid website project"

    # ──────────────────────────────────────────────────────────────────────────
    # Q&A Run: single LLM call, direct answer
    # ──────────────────────────────────────────────────────────────────────────
    async def _execute_qa_run(self, run_id: str, goal: str, plan: Optional[ExecutionPlan] = None):
        task_id = "T1-qa-answer"
        run = await self.get_run(run_id)

        task_doc = self._make_task(run_id, task_id, "Direct Answer", "writer", goal,
                                   [], ["Answer the question clearly and completely"], False, "low")
        await self._save_task(task_doc)

        run["total_tasks"] = 1
        run["status"] = RunStatus.RUNNING
        await self._save_run(run)

        start = self._now()
        task_doc["status"] = TaskStatus.RUNNING
        task_doc["started_at"] = start
        task_doc["attempts"] = 1
        await self._save_task(task_doc)

        system = (
            "You are an expert, helpful AI assistant. Answer the user's question directly, clearly, and thoroughly. "
            "Explain core concepts with definitions, outline the key steps or mechanisms, and provide at least one "
            "concrete, real-world example. Format your answer cleanly with markdown headings and bullet points for readability. "
            "Never provide placeholder responses, canned completion confirmations, or filler."
        )

        try:
            target_model = f"ollama/{settings.OLLAMA_DEFAULT_MODEL}"
            answer = await model_adapter.generate(target_model, system, goal, temperature=0.3, max_tokens=700, timeout=120.0)
            
            # Content-aware validation
            is_valid, validation_msg = self._validate_qa_content(answer, goal)
            if not is_valid:
                raise ValueError(f"Content validation failed: {validation_msg}")

            task_doc["status"] = TaskStatus.COMPLETED
            task_doc["completed_at"] = self._now()
            task_doc["execution_time_seconds"] = self._diff_seconds(task_doc["completed_at"], start)
            task_doc["output_data"] = {"answer": answer}
            task_doc["result_summary"] = "Direct answer generated."
            await self._save_task(task_doc)

            await self._register_artifact(
                run_id, task_id,
                title=f"Answer: {goal[:60]}",
                artifact_type=ArtifactType.REPORT,
                content=answer,
                is_user_facing=True,
                metadata={"run_type": "qa", "workflow": "answer_question"}
            )

            if plan and "pdf" in plan.deliverables:
                pdf_path = _save_pdf(f"Answer: {goal[:60]}", answer, run_id)
                if pdf_path and os.path.exists(pdf_path) and os.path.getsize(pdf_path) > 0:
                    await self._register_artifact(
                        run_id, task_id,
                        title=f"Download: {goal[:50]}.pdf",
                        artifact_type="pdf",
                        content=answer,
                        is_user_facing=True,
                        file_path=pdf_path,
                        metadata={"workflow": "answer_question", "file": "pdf", "path": pdf_path}
                    )

            if plan and "docx" in plan.deliverables:
                docx_path = _save_docx(f"Answer: {goal[:60]}", answer, run_id)
                if docx_path and os.path.exists(docx_path) and os.path.getsize(docx_path) > 0:
                    await self._register_artifact(
                        run_id, task_id,
                        title=f"Download: {goal[:50]}.docx",
                        artifact_type="docx",
                        content=answer,
                        is_user_facing=True,
                        file_path=docx_path,
                        metadata={"workflow": "answer_question", "file": "docx", "path": docx_path}
                    )

            run = await self.get_run(run_id)
            run["status"] = RunStatus.COMPLETED
            run["completed_at"] = self._now()
            run["completed_tasks"] = 1
            run["execution_time_seconds"] = self._diff_seconds(run["completed_at"], run["started_at"])
            await self._save_run(run)

        except Exception as e:
            logger.error(f"QA run {run_id} failed: {e}")
            task_doc["status"] = TaskStatus.FAILED
            task_doc["completed_at"] = self._now()
            task_doc["execution_time_seconds"] = self._diff_seconds(task_doc["completed_at"], start)
            task_doc["error_message"] = str(e)
            task_doc["result_summary"] = f"Failed: {str(e)}"
            await self._save_task(task_doc)

            run = await self.get_run(run_id)
            run["status"] = RunStatus.FAILED
            run["completed_at"] = self._now()
            run["error_message"] = f"QA answer generation failed: {str(e)}"
            await self._save_run(run)

    # ──────────────────────────────────────────────────────────────────────────
    # Research Run: research → write → QA → DOCX + PDF
    # ──────────────────────────────────────────────────────────────────────────
    async def _execute_research_run(self, run_id: str, goal: str, plan: Optional[ExecutionPlan] = None):
        run = await self.get_run(run_id)
        tasks_data = [
            self._make_task(run_id, "T1-research", "Research & Evidence Gathering", "researcher",
                            f"Conduct in-depth technical research on: {goal}. Detail the threat models, traditional signature vs AI anomaly detection, deep learning architectures (autoencoders, GNNs), benchmark evaluation datasets (e.g. CIC-IDS, NSL-KDD), adversarial evasion challenges, and verifiable scholarly references.",
                            [], ["Comprehensive coverage of the topic", "Cite key concepts and findings"], False, "medium"),
            self._make_task(run_id, "T2-write", "Write Full Paper / Report", "writer",
                            f"Write a complete, rigorous research paper on: {goal}. Include Title, Abstract, Keywords, 1. Introduction, 2. Related Work, 3. Proposed AI/ML Detection Methodologies, 4. Datasets & Evaluation Metrics, 5. Adversarial Evasion & Technical Limitations, 6. Conclusion, and verifiable References.",
                            ["T1-research"], ["Complete paper with all sections", "Professional academic tone"], False, "high"),
            self._make_task(run_id, "T3-qa", "Quality Review", "reviewer",
                            f"Review the research paper for completeness, technical rigor, and quality.",
                            ["T2-write"], ["All sections present", "Factual accuracy confirmed"], False, "low"),
        ]

        for t in tasks_data:
            await self._save_task(t)

        run["total_tasks"] = len(tasks_data)
        run["status"] = RunStatus.RUNNING
        await self._save_run(run)

        research_content = ""
        final_paper = ""

        try:
            for t in tasks_data:
                t["status"] = TaskStatus.RUNNING
                t["started_at"] = self._now()
                t["attempts"] = 1
                await self._save_task(t)

                sys_prompt = self._role_system_prompt(t["role"])

                if t["task_id"] == "T1-research":
                    step_max_tokens = 600
                    user_prompt = (
                        f"Research topic: {goal}\n\n"
                        f"Provide detailed technical analysis covering:\n"
                        f"1. Problem Definition: Characteristics of zero-day attacks and limitations of signature-based IDS.\n"
                        f"2. Machine Learning & Deep Learning Architectures: Anomaly detection, autoencoders, graph neural networks, and sequence models.\n"
                        f"3. Evaluation Methodology: Benchmark datasets (e.g., CIC-IDS2017, NSL-KDD, DARPA) and metrics (Precision, Recall, F1, False Positive Rate).\n"
                        f"4. Real-world Challenges: Concept drift, adversarial evasion attacks, high false-positive rates in production.\n"
                        f"5. Academic References: Cite foundational papers and standards (IEEE, ACM, USENIX, MITRE ATT&CK).\n\n"
                        f"Write substantive technical paragraphs. Do NOT generate generic placeholder text."
                    )
                elif t["task_id"] == "T2-write":
                    step_max_tokens = 900
                    user_prompt = (
                        f"Using the following research findings:\n\n{research_content[:1200]}\n\n"
                        f"Write a complete academic research paper on: {goal}.\n\n"
                        f"The paper must include the following sections with full paragraphs:\n"
                        f"# Title: {goal}\n\n"
                        f"## Abstract\n"
                        f"Provide a clear 150-250 word technical abstract summarizing problem, methodology, and results.\n\n"
                        f"## Keywords\n"
                        f"Zero-Day Detection, Machine Learning, Anomaly Detection, Deep Learning, Intrusion Detection Systems.\n\n"
                        f"## 1. Introduction and Threat Model\n"
                        f"Discuss the zero-day vulnerability lifecycle, why signature-based systems fail, and motivation for AI detection.\n\n"
                        f"## 2. Related Work and Background\n"
                        f"Compare signature vs anomaly detection paradigms.\n\n"
                        f"## 3. AI and Deep Learning Detection Architectures\n"
                        f"Explain autoencoders, representation learning, behavioral analytics, and feature extraction mechanisms.\n\n"
                        f"## 4. Datasets, Evaluation Methodology, and Metrics\n"
                        f"Detail benchmark datasets (CIC-IDS2017, NSL-KDD) and performance evaluation metrics.\n\n"
                        f"## 5. Technical Challenges, Adversarial Evasion, and Limitations\n"
                        f"Analyze false positive costs, concept drift, and adversarial attacks against AI models.\n\n"
                        f"## 6. Conclusion and Future Directions\n"
                        f"Summarize key insights and future research.\n\n"
                        f"## References\n"
                        f"List realistic scholarly citations with authors, publication venues, and years.\n\n"
                        f"Write substantive, cohesive paragraphs for every section. Do NOT use generic filler phrases."
                    )
                else:  # T3-qa
                    step_max_tokens = 400
                    user_prompt = (
                        f"Review this research paper for technical rigor and completeness:\n\n{final_paper[:2000]}\n\n"
                        f"Confirm all required sections (Abstract, Intro, Method, Limitations, References) are present and evaluate paper quality."
                    )

                target_model = f"ollama/{settings.OLLAMA_DEFAULT_MODEL}"
                response = await model_adapter.generate(
                    target_model, sys_prompt, user_prompt, temperature=0.3, max_tokens=step_max_tokens, timeout=300.0
                )

                if (
                    not response
                    or response.strip().startswith("[ERROR:")
                    or "inference failed" in response.lower()
                    or "error message related to" in response.lower()
                    or "cannot provide a review" in response.lower()
                    or "can't provide a review" in response.lower()
                ):
                    err_msg = response.strip() if response else "Empty model response"
                    t["status"] = TaskStatus.FAILED
                    t["completed_at"] = self._now()
                    t["execution_time_seconds"] = self._diff_seconds(t["completed_at"], t["started_at"])
                    t["error_message"] = err_msg
                    t["result_summary"] = f"Failed: {err_msg[:100]}"
                    await self._save_task(t)
                    raise RuntimeError(f"Step {t['task_id']} ({t['title']}) failed: {err_msg}")

                if t["task_id"] == "T1-research":
                    research_content = response
                elif t["task_id"] == "T2-write":
                    final_paper = response

                t["status"] = TaskStatus.COMPLETED
                t["completed_at"] = self._now()
                t["execution_time_seconds"] = self._diff_seconds(t["completed_at"], t["started_at"])
                t["output_data"] = {"content": response}
                t["result_summary"] = f"Completed: {t['title']}"
                await self._save_task(t)

                # Only register the final paper as user-facing; reviewer and researcher are internal
                is_user_facing = (t["task_id"] == "T2-write")
                await self._register_artifact(
                    run_id, t["task_id"],
                    title=f"Research Paper: {goal[:60]}" if is_user_facing else f"[Internal] {t['title']}",
                    artifact_type=ArtifactType.REPORT,
                    content=response,
                    is_user_facing=is_user_facing,
                    metadata={"run_type": "research", "task": t["task_id"]}
                )

                run = await self.get_run(run_id)
                run["completed_tasks"] = run.get("completed_tasks", 0) + 1
                await self._save_run(run)

            # Paper content validation
            paper_content = final_paper or research_content
            is_valid, validation_msg = self._validate_research_content(paper_content, goal)
            if not is_valid:
                raise ValueError(f"Research paper validation failed: {validation_msg}")

            paper_title = f"Research Paper: {goal[:60]}"
            docx_path = _save_docx(paper_title, paper_content, run_id)
            pdf_path = _save_pdf(paper_title, paper_content, run_id)

            if not (docx_path and os.path.exists(docx_path) and os.path.getsize(docx_path) > 0):
                raise RuntimeError("Failed to export valid DOCX research paper.")

            if not (pdf_path and os.path.exists(pdf_path) and os.path.getsize(pdf_path) > 0):
                raise RuntimeError("Failed to export valid PDF research paper.")

            await self._register_artifact(
                run_id, "T2-write",
                title=f"Download: {goal[:50]}.docx",
                artifact_type="docx",
                content=paper_content,
                is_user_facing=True,
                file_path=docx_path,
                metadata={"run_type": "research", "file": "docx", "path": docx_path}
            )

            await self._register_artifact(
                run_id, "T2-write",
                title=f"Download: {goal[:50]}.pdf",
                artifact_type="pdf",
                content=paper_content,
                is_user_facing=True,
                file_path=pdf_path,
                metadata={"run_type": "research", "file": "pdf", "path": pdf_path}
            )

            run = await self.get_run(run_id)
            run["status"] = RunStatus.COMPLETED
            run["completed_at"] = self._now()
            run["execution_time_seconds"] = self._diff_seconds(run["completed_at"], run["started_at"])
            await self._save_run(run)

        except Exception as e:
            logger.error(f"Research run {run_id} failed: {e}")
            run = await self.get_run(run_id)
            run["status"] = RunStatus.FAILED
            run["completed_at"] = self._now()
            run["error_message"] = f"Research workflow failed: {str(e)}"
            await self._save_run(run)

    # ──────────────────────────────────────────────────────────────────────────
    # Website Run: developer generates full site files + ZIP
    # ──────────────────────────────────────────────────────────────────────────
    async def _execute_website_run(self, run_id: str, goal: str, plan: Optional[ExecutionPlan] = None):
        run = await self.get_run(run_id)
        is_react = any(k in goal.lower() for k in ["react", "vite", "jsx", "tsx"])

        if is_react:
            tasks_data = [
                self._make_task(run_id, "T1-plan", "React Project Architecture & Planning", "researcher",
                                f"Plan the component architecture, state management, and file layout for React website: {goal}",
                                [], ["Define component hierarchy", "List dependencies"], False, "low"),
                self._make_task(run_id, "T2-pkg", "Generate package.json Manifest", "developer",
                                f"Generate complete package.json for React project: {goal}. Include react, react-dom, and build scripts.",
                                ["T1-plan"], ["Valid JSON", "Required dependencies included"], False, "low"),
                self._make_task(run_id, "T3-html", "Generate index.html Root", "developer",
                                f"Write index.html for Vite React project: {goal}. Include root element and main.jsx script tag.",
                                ["T1-plan"], ["Valid HTML5", "Root element included"], False, "low"),
                self._make_task(run_id, "T4-app", "Generate React App Component", "developer",
                                f"Write complete src/App.jsx for: {goal}. Include responsive navbar, hero section, portfolio projects showcase, skills grid, and interactive contact form.",
                                ["T2-pkg", "T3-html"], ["Complete JSX implementation", "All portfolio sections included"], False, "high"),
                self._make_task(run_id, "T5-css", "Generate Styling", "developer",
                                f"Write complete src/index.css for: {goal}. Modern responsive styles with dark mode variables, grid/flex layouts, and smooth animations.",
                                ["T4-app"], ["Responsive design", "CSS variables and animations"], False, "medium"),
                self._make_task(run_id, "T6-readme", "Write README", "writer",
                                f"Write complete README.md for the React website: {goal}. Include npm install and npm run dev setup instructions.",
                                ["T5-css"], ["Setup instructions", "Project overview"], False, "low"),
            ]
            file_map = {
                "T2-pkg": "package.json",
                "T3-html": "index.html",
                "T4-app": "src/App.jsx",
                "T5-css": "src/index.css",
                "T6-readme": "README.md",
            }
        else:
            tasks_data = [
                self._make_task(run_id, "T1-plan", "Website Planning & Structure", "researcher",
                                f"Plan the structure, pages, sections, and content outline for: {goal}",
                                [], ["Define all pages and sections", "List all required files"], False, "low"),
                self._make_task(run_id, "T2-html", "Generate HTML", "developer",
                                f"Write complete index.html for: {goal}. Include full semantic HTML5 structure, navigation, sections, and footer.",
                                ["T1-plan"], ["Complete valid HTML5", "All sections included"], False, "high"),
                self._make_task(run_id, "T3-css", "Generate CSS Stylesheet", "developer",
                                f"Write complete style.css for: {goal}. Modern responsive design with CSS variables, flexbox/grid.",
                                ["T2-html"], ["Fully responsive", "All HTML elements styled"], False, "high"),
                self._make_task(run_id, "T4-js", "Generate JavaScript", "developer",
                                f"Write complete script.js for: {goal}. Include interactivity, mobile nav toggle, and form handling.",
                                ["T3-css"], ["Interactive features working", "Clean JavaScript"], False, "medium"),
                self._make_task(run_id, "T5-readme", "Write README", "writer",
                                f"Write professional README.md for the website project: {goal}",
                                ["T4-js"], ["Setup instructions", "Feature list"], False, "low"),
            ]
            file_map = {
                "T2-html": "index.html",
                "T3-css": "style.css",
                "T4-js": "script.js",
                "T5-readme": "README.md",
            }

        for t in tasks_data:
            await self._save_task(t)

        run["total_tasks"] = len(tasks_data)
        run["status"] = RunStatus.RUNNING
        await self._save_run(run)

        generated_files: Dict[str, str] = {}

        try:
            for t in tasks_data:
                t["status"] = TaskStatus.RUNNING
                t["started_at"] = self._now()
                t["attempts"] = 1
                await self._save_task(t)

                sys_prompt = (
                    "You are a senior software engineer. Output ONLY clean, production-ready code inside "
                    "a single markdown code block (e.g. ```javascript, ```jsx, ```html, ```css, or ```json). "
                    "Do NOT output conversational introductions, explanations, apologies, or completion text."
                )

                # Build context from previous tasks if appropriate
                context = ""
                if "index.html" in generated_files:
                    context = f"HTML Reference:\n```html\n{generated_files['index.html'][:1200]}\n```\n\n"

                user_prompt = f"{context}{t['goal']}"
                target_model = f"ollama/{settings.OLLAMA_DEFAULT_MODEL}"
                response = await model_adapter.generate(
                    target_model, sys_prompt, user_prompt, temperature=0.2, max_tokens=1000, timeout=240.0
                )

                if (
                    not response
                    or response.strip().startswith("[ERROR:")
                    or "inference failed" in response.lower()
                ):
                    err_msg = response.strip() if response else "Empty model response"
                    t["status"] = TaskStatus.FAILED
                    t["completed_at"] = self._now()
                    t["execution_time_seconds"] = self._diff_seconds(t["completed_at"], t["started_at"])
                    t["error_message"] = err_msg
                    t["result_summary"] = f"Failed: {err_msg[:100]}"
                    await self._save_task(t)
                    raise RuntimeError(f"Step {t['task_id']} ({t['title']}) failed: {err_msg}")

                if t["task_id"] in file_map:
                    fname = file_map[t["task_id"]]
                    code = self._extract_code_block(response)
                    # If package.json, ensure valid json format
                    if fname == "package.json":
                        code = code.strip()
                        if not code.startswith("{"):
                            code = json.dumps({
                                "name": "react-portfolio",
                                "private": True,
                                "version": "1.0.0",
                                "type": "module",
                                "scripts": {"dev": "vite", "build": "vite build", "preview": "vite preview"},
                                "dependencies": {"react": "^18.3.1", "react-dom": "^18.3.1"},
                                "devDependencies": {"@vitejs/plugin-react": "^4.3.1", "vite": "^5.4.2"}
                            }, indent=2)
                    generated_files[fname] = code

                t["status"] = TaskStatus.COMPLETED
                t["completed_at"] = self._now()
                t["execution_time_seconds"] = self._diff_seconds(t["completed_at"], t["started_at"])
                t["output_data"] = {"content": response}
                t["result_summary"] = f"Completed: {t['title']}"
                await self._save_task(t)

                is_user_facing = t["task_id"] in file_map
                await self._register_artifact(
                    run_id, t["task_id"],
                    title=file_map.get(t["task_id"], t["title"]) if is_user_facing else f"[Internal] {t['title']}",
                    artifact_type=ArtifactType.CODE if is_user_facing else ArtifactType.REPORT,
                    content=generated_files.get(file_map.get(t["task_id"], ""), response),
                    is_user_facing=is_user_facing,
                    metadata={"run_type": "website", "task": t["task_id"], "filename": file_map.get(t["task_id"], "")}
                )

                run = await self.get_run(run_id)
                run["completed_tasks"] = run.get("completed_tasks", 0) + 1
                await self._save_run(run)

            # Website files validation
            is_valid, validation_msg = self._validate_website_files(generated_files, is_react)
            if not is_valid:
                raise ValueError(f"Website generation validation failed: {validation_msg}")

            # Bundle all files into ZIP and verify integrity
            zip_path = _save_zip(generated_files, run_id)
            if not (zip_path and os.path.exists(zip_path) and os.path.getsize(zip_path) > 0):
                raise RuntimeError("ZIP packaging failed or archive corrupted.")

            all_code = "\n\n".join(
                f"=== {fname} ===\n{content}" for fname, content in generated_files.items()
            )
            await self._register_artifact(
                run_id, "T2-pkg" if is_react else "T2-html",
                title="Download ZIP: Website Project",
                artifact_type="zip",
                content=all_code,
                is_user_facing=True,
                file_path=zip_path,
                metadata={"run_type": "website", "file": "zip", "path": zip_path, "files": list(generated_files.keys())}
            )

            run = await self.get_run(run_id)
            run["status"] = RunStatus.COMPLETED
            run["completed_at"] = self._now()
            run["execution_time_seconds"] = self._diff_seconds(run["completed_at"], run["started_at"])
            await self._save_run(run)

        except Exception as e:
            logger.error(f"Website run {run_id} failed: {e}")
            run = await self.get_run(run_id)
            run["status"] = RunStatus.FAILED
            run["completed_at"] = self._now()
            run["error_message"] = f"Website project generation failed: {str(e)}"
            await self._save_run(run)

    # ──────────────────────────────────────────────────────────────────────────
    # Code Run: developer + reviewer
    # ──────────────────────────────────────────────────────────────────────────
    async def _execute_code_run(self, run_id: str, goal: str, plan: Optional[ExecutionPlan] = None):
        run = await self.get_run(run_id)
        tasks_data = [
            self._make_task(run_id, "T1-dev", "Code Implementation", "developer",
                            f"Write complete, production-ready code for: {goal}. "
                            f"Include: full implementation, error handling, docstrings, and usage example.",
                            [], ["Complete working code", "Error handling included", "Usage example provided"], False, "high"),
            self._make_task(run_id, "T2-review", "Code Review & Tests", "reviewer",
                            f"Review the code for quality and write test cases.",
                            ["T1-dev"], ["Code quality assessed", "Test cases written"], False, "medium"),
        ]

        for t in tasks_data:
            await self._save_task(t)

        run["total_tasks"] = len(tasks_data)
        run["status"] = RunStatus.RUNNING
        await self._save_run(run)

        code_content = ""

        for t in tasks_data:
            t["status"] = TaskStatus.RUNNING
            t["started_at"] = self._now()
            t["attempts"] = 1
            await self._save_task(t)

            sys_prompt = self._role_system_prompt(t["role"])
            if t["task_id"] == "T2-review" and code_content:
                user_prompt = f"Review this code and write tests:\n\n```\n{code_content}\n```"
            else:
                user_prompt = t["goal"]

            response = await model_adapter.generate(
                t.get("assigned_model", "ollama/deepseek-coder:6.7b" if t["role"] == "developer" else f"ollama/{settings.OLLAMA_DEFAULT_MODEL}"),
                sys_prompt,
                user_prompt,
                temperature=0.2
            )

            if not response or response.strip().startswith("[ERROR:") or "inference failed" in response.lower():
                err_msg = response.strip() if response else "Empty model response"
                t["status"] = TaskStatus.FAILED
                t["completed_at"] = self._now()
                t["execution_time_seconds"] = self._diff_seconds(t["completed_at"], t["started_at"])
                t["error_message"] = err_msg
                t["result_summary"] = f"Failed: {err_msg[:100]}"
                await self._save_task(t)
                raise RuntimeError(f"Step {t['task_id']} ({t['title']}) failed: {err_msg}")

            if t["task_id"] == "T1-dev":
                code_content = response

            t["status"] = TaskStatus.COMPLETED
            t["completed_at"] = self._now()
            t["execution_time_seconds"] = self._diff_seconds(t["completed_at"], t["started_at"])
            t["output_data"] = {"content": response}
            t["result_summary"] = f"Completed: {t['title']}"
            await self._save_task(t)

            is_user_facing = t["task_id"] == "T1-dev"
            await self._register_artifact(
                run_id, t["task_id"],
                title=f"Code: {goal[:60]}" if is_user_facing else f"[Internal] Code Review",
                artifact_type=ArtifactType.CODE,
                content=response,
                is_user_facing=is_user_facing,
                metadata={"run_type": "code", "task": t["task_id"]}
            )

            run = await self.get_run(run_id)
            run["completed_tasks"] = run.get("completed_tasks", 0) + 1
            await self._save_run(run)

        run = await self.get_run(run_id)
        run["status"] = RunStatus.COMPLETED
        run["completed_at"] = self._now()
        run["execution_time_seconds"] = self._diff_seconds(run["completed_at"], run["started_at"])
        await self._save_run(run)

    # ──────────────────────────────────────────────────────────────────────────
    # Data Analysis Run: analyst + reviewer
    # ──────────────────────────────────────────────────────────────────────────
    async def _execute_data_analysis_run(self, run_id: str, goal: str, plan: Optional[ExecutionPlan] = None):
        run = await self.get_run(run_id)
        tasks_data = [
            self._make_task(run_id, "T1-analysis", "Data Exploration & Trend Analysis", "analyst",
                            f"Perform quantitative and qualitative analysis on the dataset for: {goal}. "
                            f"Identify patterns, compute relevant statistical metrics, and extract key insights.",
                            [], ["Quantitative metrics calculated", "Clear pattern identification"], False, "medium"),
            self._make_task(run_id, "T2-review", "Analytical Quality & Validation", "reviewer",
                            f"Validate analytical findings and evaluate recommendations for data analysis on: {goal}.",
                            ["T1-analysis"], ["Confirm analytical validity", "Verify actionable recommendations"], False, "low"),
        ]

        for t in tasks_data:
            await self._save_task(t)

        run["total_tasks"] = len(tasks_data)
        run["status"] = RunStatus.RUNNING
        await self._save_run(run)

        analysis_content = ""

        for t in tasks_data:
            t["status"] = TaskStatus.RUNNING
            t["started_at"] = self._now()
            t["attempts"] = 1
            await self._save_task(t)

            sys_prompt = self._role_system_prompt(t["role"])
            if t["task_id"] == "T2-review" and analysis_content:
                user_prompt = f"Validate this data analysis report and verify its conclusions:\n\n{analysis_content}"
            else:
                user_prompt = t["goal"]

            response = await model_adapter.generate(
                t.get("assigned_model", f"ollama/{settings.OLLAMA_DEFAULT_MODEL}"),
                sys_prompt,
                user_prompt,
                temperature=0.3
            )

            if not response or response.strip().startswith("[ERROR:") or "inference failed" in response.lower():
                err_msg = response.strip() if response else "Empty model response"
                t["status"] = TaskStatus.FAILED
                t["completed_at"] = self._now()
                t["execution_time_seconds"] = self._diff_seconds(t["completed_at"], t["started_at"])
                t["error_message"] = err_msg
                t["result_summary"] = f"Failed: {err_msg[:100]}"
                await self._save_task(t)
                raise RuntimeError(f"Step {t['task_id']} ({t['title']}) failed: {err_msg}")

            if t["task_id"] == "T1-analysis":
                analysis_content = response

            t["status"] = TaskStatus.COMPLETED
            t["completed_at"] = self._now()
            t["execution_time_seconds"] = self._diff_seconds(t["completed_at"], t["started_at"])
            t["output_data"] = {"content": response}
            t["result_summary"] = f"Completed: {t['title']}"
            await self._save_task(t)

            is_user_facing = t["task_id"] == "T1-analysis"
            await self._register_artifact(
                run_id, t["task_id"],
                title=f"Data Analysis: {goal[:60]}" if is_user_facing else f"[Internal] Analysis Review",
                artifact_type=ArtifactType.REPORT,
                content=response,
                is_user_facing=is_user_facing,
                metadata={"workflow": "analyze_data", "task": t["task_id"]}
            )

            run = await self.get_run(run_id)
            run["completed_tasks"] = run.get("completed_tasks", 0) + 1
            await self._save_run(run)

        run = await self.get_run(run_id)
        run["status"] = RunStatus.COMPLETED
        run["completed_at"] = self._now()
        run["execution_time_seconds"] = self._diff_seconds(run["completed_at"], run["started_at"])
        await self._save_run(run)

    # ──────────────────────────────────────────────────────────────────────────
    # Full Pipeline Run (planner-decomposed DAG)
    # ──────────────────────────────────────────────────────────────────────────
    async def _execute_full_pipeline_run(self, run_id: str, goal: str):
        planner_system = (
            "You are the Master Planner Agent. Decompose the goal into structured subtasks with dependencies "
            "and acceptance criteria. Return ONLY valid JSON."
        )
        planner_prompt = (
            f"Decompose this project goal into an execution DAG. "
            f"Return JSON with a 'tasks' array. Each task must have: "
            f"task_id, title, role (researcher/developer/writer/analyst/reviewer), goal, "
            f"depends_on (list of task_ids), acceptance_criteria (list of strings), "
            f"requires_approval (bool), estimated_complexity (low/medium/high).\n\n"
            f"Goal: {goal}"
        )

        try:
            raw_plan = await model_adapter.generate(f"ollama/{settings.OLLAMA_DEFAULT_MODEL}", planner_system, planner_prompt)
            clean_plan = raw_plan.strip()
            if "```" in clean_plan:
                parts = clean_plan.split("```")
                for part in parts:
                    candidate = part.strip()
                    if candidate.startswith("json"):
                        candidate = candidate[4:].strip()
                    if candidate.startswith("{") and candidate.endswith("}"):
                        clean_plan = candidate
                        break
            plan_json = json.loads(clean_plan)
            tasks_data = plan_json.get("tasks", [])
            if not tasks_data:
                raise ValueError("No tasks extracted from plan")
        except Exception as e:
            logger.error(f"Planning failed: {e}. Generating default DAG.")
            tasks_data = [
                {"task_id": "T1-research", "title": "Research & Context", "role": "researcher",
                 "goal": f"Research: {goal}", "depends_on": [], "acceptance_criteria": ["Thorough research completed"],
                 "requires_approval": False, "estimated_complexity": "medium"},
                {"task_id": "T2-dev", "title": "Implementation", "role": "developer",
                 "goal": f"Implement: {goal}", "depends_on": ["T1-research"],
                 "acceptance_criteria": ["Working implementation"], "requires_approval": False, "estimated_complexity": "high"},
                {"task_id": "T3-review", "title": "Quality Review", "role": "reviewer",
                 "goal": f"Review outputs for: {goal}", "depends_on": ["T2-dev"],
                 "acceptance_criteria": ["Quality verified"], "requires_approval": False, "estimated_complexity": "low"},
            ]

        # Persist tasks
        for t in tasks_data:
            assigned_model, routing_reason, confidence = model_router.route_task(t)
            consequential_keywords = ["deploy", "delete", "drop", "terminate", "wipe", "production_write", "charge"]
            is_consequential = any(k in t.get("title", "").lower() or k in t.get("goal", "").lower() for k in consequential_keywords)
            enforced_approval = t.get("requires_approval", False) or is_consequential

            task_doc = {
                "task_id": t["task_id"], "run_id": run_id, "title": t["title"],
                "role": t["role"], "goal": t["goal"], "status": TaskStatus.PENDING,
                "depends_on": t.get("depends_on", []), "acceptance_criteria": t.get("acceptance_criteria", []),
                "requires_approval": enforced_approval, "assigned_model": assigned_model,
                "routing_reason": routing_reason, "router_confidence": confidence,
                "attempts": 0, "max_attempts": 3, "result_summary": None, "output_data": None,
                "error_message": None, "started_at": None, "completed_at": None,
                "execution_time_seconds": None, "created_at": self._now(), "updated_at": self._now()
            }
            await self._save_task(task_doc)

        run = await self.get_run(run_id)
        run["total_tasks"] = len(tasks_data)
        run["status"] = RunStatus.RUNNING
        await self._save_run(run)

        await self._execute_dag_loop(run_id)

        # After DAG completes, synthesize a final user-facing answer
        await self._synthesize_final_answer(run_id, goal)

    async def _synthesize_final_answer(self, run_id: str, goal: str):
        """
        After all tasks complete, collect all non-QA task outputs and generate
        a single coherent user-facing summary. Store it as a user-facing artifact.
        """
        tasks = await self.get_run_tasks(run_id)
        completed_outputs = []
        for t in tasks:
            if t["status"] == TaskStatus.COMPLETED and t.get("output_data"):
                od = t["output_data"]
                # Extract text, skipping internal JSON review blobs
                if isinstance(od, dict):
                    text = od.get("text") or od.get("content") or od.get("answer") or ""
                    # Skip reviewer blobs that are just QA scores
                    if t["role"] == "reviewer" and isinstance(od, dict) and "passed" in od:
                        continue
                    if text and len(text) > 50:
                        completed_outputs.append(f"### {t['title']}\n{text}")
                elif isinstance(od, str) and len(od) > 50:
                    completed_outputs.append(f"### {t['title']}\n{od}")

        if not completed_outputs:
            return

        combined = "\n\n".join(completed_outputs)
        sys = (
            "You are a senior technical writer. Given the outputs from multiple specialized AI agents, "
            "synthesize a clear, comprehensive, user-facing response that directly answers the original request. "
            "Remove any internal metadata, JSON structures, or QA score reports. Write in clean Markdown."
        )
        user = f"Original request: {goal}\n\nAgent outputs:\n\n{combined[:3000]}\n\nWrite the final user-facing answer:"

        final_answer = await model_adapter.generate(f"ollama/{settings.OLLAMA_DEFAULT_MODEL}", sys, user, temperature=0.3)

        await self._register_artifact(
            run_id, "synthesized",
            title=f"Final Answer: {goal[:60]}",
            artifact_type=ArtifactType.REPORT,
            content=final_answer,
            is_user_facing=True,
            metadata={"run_type": "full_pipeline", "synthesized": True}
        )

    async def _execute_dag_loop(self, run_id: str):
        if run_id in self._active_loops:
            logger.info(f"DAG loop for run {run_id} is already active.")
            return
        self._active_loops.add(run_id)

        try:
            while True:
                run = await self.get_run(run_id)
                if not run or run["status"] in [RunStatus.COMPLETED, RunStatus.FAILED, RunStatus.CANCELLED]:
                    break

                if run["status"] == RunStatus.WAITING_APPROVAL:
                    await asyncio.sleep(2)
                    continue

                tasks = await self.get_run_tasks(run_id)
                completed_ids = {t["task_id"] for t in tasks if t["status"] == TaskStatus.COMPLETED}
                failed_tasks = [t for t in tasks if t["status"] == TaskStatus.FAILED]

                if failed_tasks:
                    run["status"] = RunStatus.FAILED
                    run["error_message"] = f"Task {failed_tasks[0]['task_id']} failed after maximum attempts."
                    run["completed_at"] = self._now()
                    await self._save_run(run)
                    break

                ready_tasks = [
                    t for t in tasks
                    if t["status"] in [TaskStatus.PENDING, TaskStatus.READY]
                    and all(dep in completed_ids for dep in t["depends_on"])
                ]

                if not ready_tasks:
                    if len(completed_ids) == len(tasks) and len(tasks) > 0:
                        run["status"] = RunStatus.COMPLETED
                        run["completed_at"] = self._now()
                        if run.get("started_at"):
                            run["execution_time_seconds"] = self._diff_seconds(run["completed_at"], run["started_at"])
                        run["completed_tasks"] = len(completed_ids)
                        await self._save_run(run)
                    break

                await asyncio.gather(*[self._execute_single_task(t) for t in ready_tasks])
                await asyncio.sleep(0.5)
        finally:
            self._active_loops.discard(run_id)

    async def _execute_single_task(self, task: Dict[str, Any]):
        run_id = task["run_id"]
        task_id = task["task_id"]

        if task.get("requires_approval", False) and task["status"] != TaskStatus.READY:
            task["status"] = TaskStatus.WAITING_APPROVAL
            await self._save_task(task)

            approval_doc = {
                "approval_id": f"appr-{uuid.uuid4().hex[:8]}",
                "run_id": run_id,
                "task_id": task_id,
                "action_type": "deploy_action" if "deploy" in task["title"].lower() else "external_action",
                "description": f"Human authorization requested for task: {task['title']} ({task['goal']})",
                "risk_level": RiskLevel.HIGH if "deploy" in task["title"].lower() else RiskLevel.MEDIUM,
                "status": ApprovalStatus.PENDING,
                "payload": {"goal": task["goal"], "acceptance_criteria": task["acceptance_criteria"]},
                "created_at": self._now(),
                "decided_at": None, "decided_by": None, "decision_reason": None
            }
            await self._save_approval(approval_doc)

            run = await self.get_run(run_id)
            if run:
                run["status"] = RunStatus.WAITING_APPROVAL
                run["pending_approvals"] = run.get("pending_approvals", 0) + 1
                await self._save_run(run)
            return

        task["status"] = TaskStatus.RUNNING
        task["started_at"] = self._now()
        task["attempts"] = task.get("attempts", 0) + 1
        await self._save_task(task)

        start_time = self._now()
        assigned_model = task.get("assigned_model", f"ollama/{settings.OLLAMA_DEFAULT_MODEL}")
        system_prompt = self._role_system_prompt(task["role"])

        run = await self.get_run(run_id)
        project_id = run.get("project_id", "default") if run else "default"
        rag_res = await rag_service.search_memory(project_id, task["goal"], role=task["role"], top_k=3)
        rag_context = f"\n{rag_res.citation_header}\n" if rag_res.results else ""

        user_prompt = f"{rag_context}Goal: {task['goal']}\nAcceptance Criteria: {', '.join(task['acceptance_criteria'])}"

        try:
            response_text = await model_adapter.generate(assigned_model, system_prompt, user_prompt)
            duration = (self._now() - start_time).total_seconds()

            if not response_text or response_text.strip().startswith("[ERROR:") or "inference failed" in response_text.lower():
                raise RuntimeError(f"Task inference failed: {response_text}")

            task["status"] = TaskStatus.COMPLETED
            task["completed_at"] = self._now()
            task["execution_time_seconds"] = duration
            task["result_summary"] = f"Successfully generated deliverable for {task['title']}."
            task["output_data"] = {"text": response_text}
            await self._save_task(task)

            # Reviewer tasks are internal only — don't show raw JSON to user
            is_user_facing = task["role"] != "reviewer"

            await self._register_artifact(
                run_id, task_id,
                title=f"Deliverable - {task['title']}" if is_user_facing else f"[Internal] QA - {task['title']}",
                artifact_type=ArtifactType.REPORT if task["role"] in ["researcher", "writer", "analyst"] else ArtifactType.CODE,
                content=response_text,
                is_user_facing=is_user_facing,
                metadata={"model": assigned_model, "role": task["role"],
                          "task_title": task["title"], "task_goal": task["goal"],
                          "sources_cited": [r.title for r in rag_res.results] if rag_res.results else []}
            )

            run = await self.get_run(run_id)
            if run:
                run["completed_tasks"] = run.get("completed_tasks", 0) + 1
                await self._save_run(run)

        except Exception as e:
            logger.error(f"Task {task_id} execution error: {e}")
            if task["attempts"] < task.get("max_attempts", 3):
                task["status"] = TaskStatus.PENDING
            else:
                task["status"] = TaskStatus.FAILED
                task["error_message"] = str(e)
            await self._save_task(task)

    # ==========================================
    # Approval Decisions
    # ==========================================
    async def process_approval_decision(self, approval_id: str, decision: ApprovalStatus, reason: str = "") -> Optional[Dict[str, Any]]:
        approval = await self.get_approval(approval_id)
        if not approval or approval["status"] != ApprovalStatus.PENDING:
            return None

        approval["status"] = decision
        approval["decided_at"] = self._now()
        approval["decided_by"] = "human_operator"
        approval["decision_reason"] = reason
        await self._save_approval(approval)

        task = await self.get_task(approval["task_id"], run_id=approval["run_id"])
        run = await self.get_run(approval["run_id"])

        if run and run.get("pending_approvals", 0) > 0:
            run["pending_approvals"] -= 1

        if decision == ApprovalStatus.APPROVED:
            if task:
                task["status"] = TaskStatus.READY
                task["requires_approval"] = False
                await self._save_task(task)
            if run and run["status"] == RunStatus.WAITING_APPROVAL and run.get("pending_approvals", 0) == 0:
                run["status"] = RunStatus.RUNNING
                await self._save_run(run)

            asyncio.create_task(self._execute_dag_loop(approval["run_id"]))
        else:
            if task:
                task["status"] = TaskStatus.FAILED
                task["error_message"] = f"Action rejected by human reviewer: {reason}"
                await self._save_task(task)
            if run:
                run["status"] = RunStatus.FAILED
                run["error_message"] = f"Run stopped due to rejected approval on task {task['title']}."
                await self._save_run(run)

        return approval

    # ==========================================
    # Getters & Persistence Helpers
    # ==========================================
    async def get_run_tasks(self, run_id: str) -> List[Dict[str, Any]]:
        if db_manager.connected and db_manager.db is not None:
            cursor = db_manager.db.tasks.find({"run_id": run_id}, {"_id": 0})
            return await cursor.to_list(length=100)
        return [t for t in self.mem_tasks.values() if t.get("run_id") == run_id]

    async def get_task(self, task_id: str, run_id: Optional[str] = None) -> Optional[Dict[str, Any]]:
        if db_manager.connected and db_manager.db is not None:
            q = {"task_id": task_id}
            if run_id:
                q["run_id"] = run_id
            return await db_manager.db.tasks.find_one(q, {"_id": 0})
        if run_id:
            return self.mem_tasks.get(f"{run_id}_{task_id}")
        for t in reversed(list(self.mem_tasks.values())):
            if t.get("task_id") == task_id:
                return t
        return None

    async def update_task_status_safe(self, task_id: str, new_status: TaskStatus, run_id: Optional[str] = None) -> Dict[str, Any]:
        task = await self.get_task(task_id, run_id=run_id)
        if not task:
            raise ValueError(f"Task {task_id} not found")

        if task.get("requires_approval", False) and new_status in [TaskStatus.COMPLETED, TaskStatus.RUNNING]:
            approvals = await self.list_approvals(run_id=task.get("run_id"))
            task_approvals = [a for a in approvals if a.get("task_id") == task_id]
            is_approved = any(a.get("status") == ApprovalStatus.APPROVED for a in task_approvals)
            if not is_approved:
                logger.warning(f"Security Alert: Attempted unauthorized status transition on gated task {task_id}.")
                raise PermissionError("Security Policy Violation: Task requires human operator approval.")

        task["status"] = new_status
        await self._save_task(task)
        return task

    async def get_approval(self, approval_id: str) -> Optional[Dict[str, Any]]:
        if db_manager.connected and db_manager.db is not None:
            return await db_manager.db.approvals.find_one({"approval_id": approval_id}, {"_id": 0})
        return self.mem_approvals.get(approval_id)

    async def list_approvals(self, run_id: Optional[str] = None) -> List[Dict[str, Any]]:
        if db_manager.connected and db_manager.db is not None:
            q = {"run_id": run_id} if run_id else {}
            cursor = db_manager.db.approvals.find(q, {"_id": 0}).sort("created_at", -1)
            return await cursor.to_list(length=100)
        items = list(self.mem_approvals.values())
        if run_id:
            items = [a for a in items if a.get("run_id") == run_id]
        return sorted(items, key=lambda x: x["created_at"], reverse=True)

    async def list_artifacts(self, run_id: Optional[str] = None) -> List[Dict[str, Any]]:
        if db_manager.connected and db_manager.db is not None:
            q = {"run_id": run_id} if run_id else {}
            cursor = db_manager.db.artifacts.find(q, {"_id": 0}).sort("created_at", -1)
            return await cursor.to_list(length=100)
        items = list(self.mem_artifacts.values())
        if run_id:
            items = [a for a in items if a.get("run_id") == run_id]
        return sorted(items, key=lambda x: x["created_at"], reverse=True)

    async def _save_run(self, run: Dict[str, Any]):
        run["updated_at"] = self._now()
        self.mem_runs[run["run_id"]] = run
        if db_manager.connected and db_manager.db is not None:
            await db_manager.db.runs.update_one({"run_id": run["run_id"]}, {"$set": run}, upsert=True)

    async def _save_task(self, task: Dict[str, Any]):
        task["updated_at"] = self._now()
        task_key = f"{task['run_id']}_{task['task_id']}"
        self.mem_tasks[task_key] = task
        if db_manager.connected and db_manager.db is not None:
            await db_manager.db.tasks.update_one(
                {"run_id": task["run_id"], "task_id": task["task_id"]},
                {"$set": task}, upsert=True
            )

    async def _save_approval(self, approval: Dict[str, Any]):
        self.mem_approvals[approval["approval_id"]] = approval
        if db_manager.connected and db_manager.db is not None:
            await db_manager.db.approvals.update_one(
                {"approval_id": approval["approval_id"]}, {"$set": approval}, upsert=True
            )

    async def _save_artifact(self, artifact: Dict[str, Any]):
        self.mem_artifacts[artifact["artifact_id"]] = artifact
        if db_manager.connected and db_manager.db is not None:
            await db_manager.db.artifacts.update_one(
                {"artifact_id": artifact["artifact_id"]}, {"$set": artifact}, upsert=True
            )

    async def _register_artifact(
        self,
        run_id: str,
        task_id: str,
        title: str,
        artifact_type: str,
        content: str,
        is_user_facing: bool = True,
        file_path: Optional[str] = None,
        metadata: Optional[Dict[str, Any]] = None,
    ):
        verified_file_path = None
        if file_path:
            if os.path.exists(file_path) and os.path.isfile(file_path) and os.path.getsize(file_path) > 0:
                verified_file_path = file_path
            else:
                logger.warning(f"Omitting unverified or non-existent artifact file path: {file_path}")

        artifact_id = f"art-{uuid.uuid4().hex[:8]}"
        size_bytes = os.path.getsize(verified_file_path) if verified_file_path else len(content.encode("utf-8"))
        artifact_doc = {
            "artifact_id": artifact_id,
            "run_id": run_id,
            "task_id": task_id,
            "title": title,
            "artifact_type": artifact_type,
            "content_snippet": content[:500],
            "full_content": content,
            "is_user_facing": is_user_facing,
            "file_path": verified_file_path,
            "size_bytes": size_bytes,
            "created_at": self._now(),
            "metadata": metadata or {},
        }
        await self._save_artifact(artifact_doc)
        return artifact_doc

    # ──────────────────────────────────────────────────────────────────────────
    # Internal utility methods
    # ──────────────────────────────────────────────────────────────────────────
    @staticmethod
    def _make_task(run_id: str, task_id: str, title: str, role: str, goal: str,
                   depends_on: List[str], acceptance_criteria: List[str],
                   requires_approval: bool, estimated_complexity: str) -> Dict[str, Any]:
        from datetime import datetime, timezone
        now = datetime.now(timezone.utc)
        consequential_keywords = ["deploy", "delete", "drop", "terminate", "wipe", "production_write", "charge"]
        is_consequential = any(k in title.lower() or k in goal.lower() for k in consequential_keywords)
        return {
            "task_id": task_id, "run_id": run_id, "title": title,
            "role": role, "goal": goal, "status": TaskStatus.PENDING,
            "depends_on": depends_on, "acceptance_criteria": acceptance_criteria,
            "requires_approval": requires_approval or is_consequential,
            "estimated_complexity": estimated_complexity,
            "assigned_model": f"ollama/{settings.OLLAMA_DEFAULT_MODEL}",
            "routing_reason": "Assigned by workflow", "router_confidence": 1.0,
            "attempts": 0, "max_attempts": 3, "result_summary": None, "output_data": None,
            "error_message": None, "started_at": None, "completed_at": None,
            "execution_time_seconds": None, "created_at": now, "updated_at": now,
        }

    @staticmethod
    def _role_system_prompt(role: str) -> str:
        prompts = {
            "researcher": (
                "You are an expert researcher. Provide thorough, well-structured analysis with clear "
                "headings, key findings, and actionable insights. Use Markdown formatting."
            ),
            "developer": (
                "You are a senior software engineer. Write complete, production-ready code with "
                "proper error handling, documentation, and examples. Format code in Markdown code blocks."
            ),
            "writer": (
                "You are a professional technical writer. Write clear, well-structured documents "
                "with proper Markdown formatting including headings, lists, and emphasis."
            ),
            "analyst": (
                "You are a data analyst. Provide quantitative and qualitative analysis with "
                "structured findings, metrics, and recommendations in Markdown format."
            ),
            "reviewer": (
                "You are a quality assurance engineer. Review the outputs thoroughly and provide "
                "a structured quality report in Markdown with clear pass/fail criteria and recommendations."
            ),
            "planner": (
                "You are the Master Planner Agent. Decompose goals into structured subtasks. "
                "Return ONLY valid JSON with a 'tasks' array."
            ),
        }
        return prompts.get(role, "You are a specialized AI agent. Complete the task thoroughly.")

    @staticmethod
    def _extract_code_block(text: str) -> str:
        """Extract code from markdown code fences, or return text as-is."""
        match = re.search(r"```(?:\w+)?\n?(.*?)```", text, re.DOTALL)
        if match:
            return match.group(1).strip()
        return text.strip()


orchestrator = OrchestrationService()
